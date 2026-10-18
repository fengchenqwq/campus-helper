# Coze 智能体调用封装（基于本地知识库检索）
import jieba
from collections import OrderedDict

from cozepy import Coze, ChatEventType, Message, MessageObjectString, MessageObjectStringType, MessageType, TokenAuth

from config import COZE_API_BASE, COZE_API_TOKEN, COZE_BOT_ID
from database import get_db

# 精准提示词：引导智能体以校园助手身份回答（专业问题走知识库，闲聊自然回应）
SYSTEM_PROMPT = (
    '你是「校园百事通」校园助手。'
    '当用户咨询校园相关问题时，优先依据提供的知识库内容回答；知识库没有的，说明以学校官方通知为准，不要编造。'
    '当用户进行日常闲聊或问校园以外的问题时，自然友好地回应。'
    '用中文简洁回答。'
)

_coze = Coze(auth=TokenAuth(token=COZE_API_TOKEN), base_url=COZE_API_BASE)

# 回答缓存：问题 -> 完整回复（内存 LRU + SQLite 持久化，避免重复问题反复调 Coze）
_cache = OrderedDict()
_CACHE_MAX = 500

# 缓存持久化：服务重启后从 SQLite 加载回内存，缓存不清空
_CACHE_TABLE_SQL = '''
CREATE TABLE IF NOT EXISTS answer_cache (
    question   TEXT PRIMARY KEY,
    answer     TEXT NOT NULL,
    updated_at DATETIME NOT NULL DEFAULT CURRENT_TIMESTAMP
);
'''


def _load_cache():
    """启动时建表并把持久化缓存加载进内存（失败不影响启动）。"""
    try:
        conn = get_db()
        conn.execute(_CACHE_TABLE_SQL)
        conn.commit()
        rows = conn.execute('SELECT question, answer FROM answer_cache ORDER BY updated_at, rowid').fetchall()
        conn.close()
        for r in rows:
            _cache[r['question']] = r['answer']
    except Exception:
        pass


def _persist(key, answer):
    """缓存写入 SQLite 持久化（失败不影响主流程）。"""
    try:
        conn = get_db()
        conn.execute(
            'INSERT OR REPLACE INTO answer_cache (question, answer, updated_at) VALUES (?, ?, CURRENT_TIMESTAMP)',
            (key, answer),
        )
        conn.commit()
        conn.close()
    except Exception:
        pass


def _remove(key):
    """缓存淘汰时同步删除 SQLite 记录（失败不影响主流程）。"""
    try:
        conn = get_db()
        conn.execute('DELETE FROM answer_cache WHERE question=?', (key,))
        conn.commit()
        conn.close()
    except Exception:
        pass


# 启动时加载持久化缓存
_load_cache()


def _search_knowledge(question: str, limit: int = 3) -> str:
    """从本地知识库检索相关内容（jieba 分词 + 关键词匹配）。"""
    conn = get_db()
    rows = conn.execute('SELECT category, question, answer FROM knowledge').fetchall()
    conn.close()

    stopwords = {
        '的', '了', '吗', '呢', '是', '什么', '怎么', '如何', '怎样', '请', '帮我', '请问',
        '一下', '安排', '申请', '查询', '了解', '求', '啊', '呀', '吧', '么', '有', '在',
        '和', '与', '及', '或', '我', '你', '他', '她', '它', '们', '这', '那', '哪', '个', '些', '不', '要', '能',
    }
    keywords = [w for w in jieba.cut(question) if w.strip() and w not in stopwords and len(w) >= 2]
    if not keywords:
        return ''

    # 统计每条知识命中的关键词数：question 字段权重高于 answer，按得分排序
    scored = []
    for r in rows:
        q_text = r['question'] or ''
        a_text = r['answer'] or ''
        score = sum(2 for kw in keywords if kw in q_text) + sum(1 for kw in keywords if kw in a_text)
        if score > 0:
            scored.append((score, r))

    scored.sort(key=lambda x: -x[0])

    # 去重：相同 question 只保留得分最高的一条，避免重复记录挤占结果
    seen = set()
    uniq = []
    for score, r in scored:
        if r['question'] in seen:
            continue
        seen.add(r['question'])
        uniq.append((score, r))

    top = uniq[:limit]
    if not top:
        return ''

    parts = []
    for _, r in top:
        answer = (r['answer'] or '')[:300]  # 截断，避免提示词过长
        parts.append(f'[{r["category"]}] 问：{r["question"]} 答：{answer}')
    return '\n'.join(parts)


def chat(user_id: str, question: str) -> str:
    """调用 Coze 智能体，返回完整回答文本。"""
    knowledge = _search_knowledge(question)

    if knowledge:
        content = (
            SYSTEM_PROMPT + '\n\n'
            '以下是本地校园知识库中检索到的相关内容，请优先参考这些内容回答：\n'
            + knowledge + '\n\n'
            '用户问题：' + question
        )
    else:
        content = SYSTEM_PROMPT + '\n\n用户问题：' + question

    result = _coze.chat.create_and_poll(
        bot_id=COZE_BOT_ID,
        user_id=user_id,
        additional_messages=[Message.build_user_question_text(content)],
        auto_save_history=True,
    )
    answer = ''
    for msg in result.messages or []:
        if msg.type == MessageType.ANSWER and msg.content:
            answer += msg.content
    return answer.strip()


def recognize_image(file_bytes: bytes, filename: str) -> str:
    """上传图片到 Coze，识别图片中的文字/内容。"""
    uploaded = _coze.files.upload(file=(filename, file_bytes))
    objects = [
        MessageObjectString(
            type=MessageObjectStringType.TEXT,
            text='请识别这张图片中的文字内容并原样输出；如果图中没有文字，请简要描述图片内容。',
        ),
        MessageObjectString(type=MessageObjectStringType.IMAGE, file_id=uploaded.id),
    ]
    message = Message.build_user_question_objects(objects)
    result = _coze.chat.create_and_poll(
        bot_id=COZE_BOT_ID,
        user_id='image_recognizer',
        additional_messages=[message],
        auto_save_history=True,
    )
    answer = ''
    for msg in result.messages or []:
        if msg.type == MessageType.ANSWER and msg.content:
            answer += msg.content
    return answer.strip()


def parse_document(file_bytes: bytes, filename: str) -> str:
    """解析文档（docx/pdf/pptx/xlsx），提取文字内容。"""
    import io
    ext = filename.lower().rsplit('.', 1)[-1] if '.' in filename else ''
    try:
        if ext == 'doc':
            import subprocess
            import tempfile
            import os
            with tempfile.NamedTemporaryFile(suffix='.doc', delete=False) as f:
                f.write(file_bytes)
                tmp = f.name
            try:
                result = subprocess.run(['antiword', tmp], capture_output=True, text=True, timeout=30)
                return result.stdout.strip()
            finally:
                os.unlink(tmp)
        if ext == 'docx':
            import docx
            doc = docx.Document(io.BytesIO(file_bytes))
            return '\n'.join(p.text for p in doc.paragraphs if p.text.strip())
        if ext == 'pdf':
            import fitz
            doc = fitz.open(stream=file_bytes, filetype='pdf')
            return '\n'.join(page.get_text() for page in doc)
        if ext == 'pptx':
            import pptx
            prs = pptx.Presentation(io.BytesIO(file_bytes))
            parts = []
            for slide in prs.slides:
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        parts.append(shape.text_frame.text)
            return '\n'.join(parts)
        if ext in ('xlsx', 'xls'):
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True)
            parts = []
            for ws in wb.worksheets:
                for row in ws.iter_rows(values_only=True):
                    parts.append(' '.join(str(c) for c in row if c is not None))
            return '\n'.join(parts)
        return ''
    except Exception:
        return ''


def chat_stream(user_id: str, question: str):
    """流式调用 Coze，逐条 yield 事件字典 {type: delta/done/error, content}。"""
    key = question.strip()

    # 缓存命中：直接返回缓存回复，不调 Coze
    if key in _cache:
        _cache.move_to_end(key)
        yield {'type': 'delta', 'content': _cache[key]}
        yield {'type': 'done', 'content': ''}
        return

    knowledge = _search_knowledge(question)
    if knowledge:
        content = (
            SYSTEM_PROMPT + '\n\n'
            '以下是本地校园知识库中检索到的相关内容，请优先参考这些内容回答：\n'
            + knowledge + '\n\n'
            '用户问题：' + question
        )
    else:
        content = SYSTEM_PROMPT + '\n\n用户问题：' + question

    full_reply = ''
    has_error = False
    try:
        for event in _coze.chat.stream(
            bot_id=COZE_BOT_ID,
            user_id=user_id,
            additional_messages=[Message.build_user_question_text(content)],
            auto_save_history=True,
        ):
            if event.event == ChatEventType.CONVERSATION_MESSAGE_DELTA:
                msg = event.message
                if msg and msg.type == MessageType.ANSWER and msg.content:
                    full_reply += msg.content
                    yield {'type': 'delta', 'content': msg.content}
            elif event.event == ChatEventType.CONVERSATION_CHAT_COMPLETED:
                yield {'type': 'done', 'content': ''}
            elif event.event == ChatEventType.CONVERSATION_CHAT_FAILED:
                has_error = True
                err = getattr(getattr(event, 'chat', None), 'last_error', None)
                text = getattr(err, 'msg', '') or '对话失败'
                yield {'type': 'error', 'content': text}
    except Exception as exc:  # noqa: BLE001
        has_error = True
        yield {'type': 'error', 'content': str(exc)}
    finally:
        # 成功后缓存完整回复（内存 LRU + SQLite 持久化）
        if full_reply and not has_error:
            _cache[key] = full_reply
            _persist(key, full_reply)
            if len(_cache) > _CACHE_MAX:
                evicted_key, _ = _cache.popitem(last=False)
                _remove(evicted_key)


def summarize_title(dialog_text: str) -> str:
    """把对话内容总结成 20 字左右的标题。"""
    prompt = (
        '请把以下对话内容总结成一个不超过20字的标题，概括对话核心主题。'
        '直接输出标题本身，不要加引号、标点或任何解释：\n\n' + dialog_text
    )
    result = _coze.chat.create_and_poll(
        bot_id=COZE_BOT_ID,
        user_id='summarizer',
        additional_messages=[Message.build_user_question_text(prompt)],
        auto_save_history=True,
    )
    answer = ''
    for msg in result.messages or []:
        if msg.type == MessageType.ANSWER and msg.content:
            answer += msg.content
    return answer.strip()[:30]

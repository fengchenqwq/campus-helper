# 知识库批量导入脚本 v2：解析「校园文件」文件夹中的各类文档并写入 knowledge 表
import os
import shutil
import sqlite3
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path

import fitz  # pymupdf
import openpyxl
import pytesseract
import win32com.client
import xlrd
from docx import Document
from PIL import Image
from pptx import Presentation

FOLDER = Path('C:/Users/14218/Desktop/校园文件')
DB_PATH = Path(__file__).parent / 'star_dream.db'
TESSDATA = str(Path(__file__).parent / 'tessdata')

pytesseract.pytesseract.tesseract_cmd = r'C:\Program Files\Tesseract-OCR\tesseract.exe'
os.environ['TESSDATA_PREFIX'] = TESSDATA

SEVEN_ZIP = r'C:\Program Files\7-Zip\7z.exe'

_word = None


def get_word():
    global _word
    if _word is None:
        _word = win32com.client.Dispatch('Word.Application')
        _word.Visible = False
    return _word


def close_word():
    global _word
    if _word is not None:
        try:
            _word.Quit()
        except Exception:
            pass
        _word = None


# ===== 各格式文本提取 =====
def extract_docx(path):
    doc = Document(path)
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    parts.append(cell.text.strip())
    return '\n'.join(parts)


def extract_pdf(path):
    doc = fitz.open(path)
    text = ''
    for page in doc:
        text += page.get_text()
    doc.close()
    if text.strip():
        return text
    # 扫描版 PDF：转图片 OCR
    return ocr_pdf(path)


def ocr_pdf(path):
    doc = fitz.open(path)
    full = ''
    for page in doc:
        pix = page.get_pixmap(dpi=200)
        img = Image.frombytes('RGB', [pix.width, pix.height], pix.samples)
        full += pytesseract.image_to_string(img, lang='chi_sim+eng') + '\n'
    doc.close()
    return full


def extract_doc(path):
    word = get_word()
    doc = word.Documents.Open(str(path))
    text = doc.Content.Text
    doc.Close(False)
    return text


def extract_pptx(path):
    prs = Presentation(path)
    parts = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                parts.append(shape.text_frame.text)
    return '\n'.join(parts)


def extract_xlsx(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    parts = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            parts.append(' '.join(str(c) for c in row if c is not None))
    return '\n'.join(parts)


def extract_xls(path):
    wb = xlrd.open_workbook(path)
    parts = []
    for ws in wb.sheets():
        for r in range(ws.nrows):
            parts.append(' '.join(str(c) for c in ws.row_values(r)))
    return '\n'.join(parts)


def classify(name):
    if any(k in name for k in ['补考', '考试', '四六级', '选课', '学分', '评教', '缴费', '学业预警', '培养', '课程', '教学', '教育', '评优', '成绩']):
        return '教务'
    if any(k in name for k in ['奖学金', '助学金', '宿舍', '公寓', '处分', '辅导员', '五好', '学风', '困难', '认定', '鉴定', '体测', '学平险', '医保', '楼长', '层长', '班委', '新生']):
        return '生活'
    if any(k in name for k in ['社团', '大赛', '创新', '活动', '书屋', '社会实践', '创业', '职业规划', '志愿', '园区', 'OPC']):
        return '社团'
    return 'FAQ'


def chunk_text(text, size=800):
    paragraphs = [p.strip() for p in text.split('\n') if p.strip()]
    chunks = []
    current = ''
    for p in paragraphs:
        if current and len(current) + len(p) > size:
            chunks.append(current)
            current = p
        else:
            current = current + '\n' + p if current else p
    if current:
        chunks.append(current)
    return chunks


# ===== 单个文件处理 =====
def process_file(f: Path, cur, inserted, skipped):
    name = f.name
    if name.startswith('~$'):
        return inserted, skipped
    suffix = f.suffix.lower()
    extractors = {
        '.docx': extract_docx,
        '.pdf': extract_pdf,
        '.doc': extract_doc,
        '.pptx': extract_pptx,
        '.xlsx': extract_xlsx,
        '.xls': extract_xls,
    }
    if suffix not in extractors:
        skipped.append((name, '不支持的格式 ' + suffix))
        return inserted, skipped

    try:
        text = extractors[suffix](f)
    except Exception as exc:  # noqa: BLE001
        skipped.append((name, '解析失败: ' + str(exc)[:80]))
        return inserted, skipped

    if not text.strip():
        skipped.append((name, '无文本内容'))
        return inserted, skipped

    category = classify(name)
    title = f.stem
    for chunk in chunk_text(text):
        cur.execute(
            'INSERT INTO knowledge (category, question, answer) VALUES (?, ?, ?)',
            (category, title, chunk),
        )
        inserted += 1
    return inserted, skipped


def main():
    sys.stdout.reconfigure(errors='replace')
    conn = sqlite3.connect(DB_PATH)
    cur = conn.cursor()
    cur.execute('DELETE FROM knowledge')

    inserted = 0
    skipped = []
    tmp_dirs = []

    try:
        for f in sorted(FOLDER.iterdir()):
            suffix = f.suffix.lower()
            if suffix in ('.zip', '.rar'):
                # 解压到临时目录，递归处理里面的文件
                tmp = tempfile.mkdtemp()
                tmp_dirs.append(tmp)
                try:
                    if suffix == '.zip':
                        with zipfile.ZipFile(f) as zf:
                            zf.extractall(tmp)
                    else:
                        subprocess.run(
                            [SEVEN_ZIP, 'x', str(f), '-o' + tmp, '-y'],
                            capture_output=True, check=True,
                        )
                except Exception as exc:  # noqa: BLE001
                    skipped.append((f.name, suffix + ' 解压失败: ' + str(exc)[:80]))
                    continue
                for inner in sorted(Path(tmp).rglob('*')):
                    if inner.is_file():
                        inserted, skipped = process_file(inner, cur, inserted, skipped)
            else:
                inserted, skipped = process_file(f, cur, inserted, skipped)
    finally:
        # 先提交数据并关闭 Word（释放 .doc 文件句柄），再清理临时目录
        conn.commit()
        conn.close()
        close_word()
        for tmp in tmp_dirs:
            shutil.rmtree(tmp, ignore_errors=True)

    print('导入完成：共插入', inserted, '条知识')
    if skipped:
        print('\n跳过/失败文件：')
        for name, reason in skipped:
            print('  -', name, '|', reason)


if __name__ == '__main__':
    main()
